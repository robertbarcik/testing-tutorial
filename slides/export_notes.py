#!/usr/bin/env python3
"""Export speaker notes from 00_intro_testing_by_hand.pptx to 00_intro_speaker_notes.md."""
from pathlib import Path
from pptx import Presentation
here = Path(__file__).resolve().parent
p = Presentation(here / "00_intro_testing_by_hand.pptx")
lines = ["# Chapters 0–1 (slides) — speaker notes", "",
         "Deck: `slides/00_intro_testing_by_hand.pptx` (per-slide JPGs in `slides-jpg/00_intro/`).",
         "Long-form reading with worked GPT-5 outputs: `00_Basic_Testing.md`.",
         "Lecture split (Skillmea sheet, 14-slide deck since 2026-08-17): 0_1 = slides 1–2 · 0_2 = 3 · 1_1 = 4–5 · 1_2 = 6–9 · 1_3 = 10–12 · 1_4 = 13 · 1_5 = 14.", ""]
for i, s in enumerate(p.slides, 1):
    best = ("", 0)
    for sh in s.shapes:
        if sh.has_text_frame:
            for para in sh.text_frame.paragraphs:
                for r in para.runs:
                    if r.font.size and r.font.size.pt > best[1]:
                        best = (sh.text_frame.text.replace("\n", " "), r.font.size.pt)
    lines += [f"## {i}. {best[0]}", "", s.notes_slide.notes_text_frame.text, ""]
(here / "00_intro_speaker_notes.md").write_text("\n".join(lines))
print("exported", len(p.slides), "slides")
